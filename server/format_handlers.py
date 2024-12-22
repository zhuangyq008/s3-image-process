import os
import io
from typing import List, Tuple, Optional
from PIL import Image
import fitz  # PyMuPDF
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table, _Row
from docx.text.paragraph import Paragraph
from pptx import Presentation
import pandas as pd
from tempfile import NamedTemporaryFile
import logging

# 设置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class FormatHandler:
    def convert_to_pdf(self, input_data: bytes, dpi: int = 300) -> bytes:
        raise NotImplementedError
        
    def convert_to_images(self, input_data: bytes, dpi: int = 300) -> List[Tuple[str, bytes]]:
        raise NotImplementedError
        
    def convert_to_text(self, input_data: bytes) -> bytes:
        raise NotImplementedError

class PDFHandler(FormatHandler):
    def convert_to_pdf(self, input_data: bytes, dpi: int = 300) -> bytes:
        return input_data  # Already PDF
        
    def convert_to_images(self, input_data: bytes, dpi: int = 300) -> List[Tuple[str, bytes]]:
        images = []
        pdf_document = None
        try:
            logger.info("Opening PDF document")
            pdf_document = fitz.open(stream=input_data, filetype="pdf")
            page_count = len(pdf_document)
            logger.info(f"Processing PDF with {page_count} pages")
            
            for page_num in range(page_count):
                logger.info(f"Processing page {page_num + 1}/{page_count}")
                page = pdf_document[page_num]
                
                # 计算合适的缩放因子
                zoom = dpi/72
                mat = fitz.Matrix(zoom, zoom)
                
                try:
                    # 获取页面图像
                    pix = page.get_pixmap(matrix=mat, alpha=False)
                    
                    # 转换为PIL Image
                    img_data = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    
                    # 保存为PNG
                    img_byte_arr = io.BytesIO()
                    img_data.save(img_byte_arr, format='PNG', optimize=True)
                    
                    images.append((f"page_{page_num + 1}.png", img_byte_arr.getvalue()))
                    
                    # 清理
                    img_byte_arr.close()
                    del pix
                    del img_data
                    
                except Exception as e:
                    logger.error(f"Error processing page {page_num + 1}: {str(e)}")
                    raise ValueError(f"Failed to convert page {page_num + 1} to image: {str(e)}")
            
            logger.info("Successfully converted all pages to images")
            return images
            
        except Exception as e:
            logger.error(f"PDF conversion error: {str(e)}")
            raise ValueError(f"Failed to convert PDF to images: {str(e)}")
            
        finally:
            if pdf_document:
                try:
                    pdf_document.close()
                except Exception as e:
                    logger.error(f"Error closing PDF document: {str(e)}")
        
    def convert_to_text(self, input_data: bytes) -> bytes:
        pdf_document = None
        try:
            pdf_document = fitz.open(stream=input_data, filetype="pdf")
            text = ""
            for page in pdf_document:
                text += page.get_text()
            return text.encode('utf-8')
        finally:
            if pdf_document:
                pdf_document.close()

class WordHandler(FormatHandler):
    def iter_block_items(self, parent):
        """
        Iterate through all paragraphs and tables in a document:
        """
        if isinstance(parent, _Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            raise ValueError("Something's not right")

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    def extract_text_from_table(self, table):
        """Extract text from a table"""
        text = []
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                # 递归处理单元格中的内容
                cell_text = []
                for block in self.iter_block_items(cell):
                    if isinstance(block, Paragraph):
                        cell_text.append(block.text.strip())
                    elif isinstance(block, Table):
                        cell_text.append(self.extract_text_from_table(block))
                row_text.append(' '.join(cell_text))
            text.append(' | '.join(filter(None, row_text)))
        return '\n'.join(text)

    def convert_to_pdf(self, input_data: bytes, dpi: int = 300) -> bytes:
        with NamedTemporaryFile(delete=False, suffix='.docx') as tmp_in, \
             NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_out:
            try:
                tmp_in.write(input_data)
                tmp_in.flush()
                
                # Use libreoffice for conversion
                cmd = f'soffice --headless --convert-to pdf "{tmp_in.name}" --outdir "{os.path.dirname(tmp_out.name)}"'
                logger.info(f"Executing command: {cmd}")
                result = os.system(cmd)
                
                if result != 0:
                    raise ValueError(f"LibreOffice conversion failed with exit code {result}")
                
                with open(tmp_out.name, 'rb') as f:
                    pdf_data = f.read()
                    
                return pdf_data
                
            finally:
                try:
                    os.unlink(tmp_in.name)
                    os.unlink(tmp_out.name)
                except Exception as e:
                    logger.error(f"Error cleaning up temporary files: {str(e)}")
            
    def convert_to_images(self, input_data: bytes, dpi: int = 300) -> List[Tuple[str, bytes]]:
        pdf_data = self.convert_to_pdf(input_data)
        pdf_handler = PDFHandler()
        return pdf_handler.convert_to_images(pdf_data, dpi)
        
    def convert_to_text(self, input_data: bytes) -> bytes:
        try:
            logger.info("Opening Word document")
            doc = Document(io.BytesIO(input_data))
            
            text_parts = []
            
            # 处理文档主体
            logger.info("Processing document content")
            for block in self.iter_block_items(doc):
                if isinstance(block, Paragraph):
                    if block.text.strip():  # 只添加非空段落
                        text_parts.append(block.text.strip())
                elif isinstance(block, Table):
                    table_text = self.extract_text_from_table(block)
                    if table_text.strip():  # 只添加非空表格
                        text_parts.append(table_text)
            
            # 处理页眉
            logger.info("Processing headers")
            for section in doc.sections:
                header = section.header
                if header:
                    for paragraph in header.paragraphs:
                        if paragraph.text.strip():
                            text_parts.append(f"[Header] {paragraph.text.strip()}")
            
            # 处理页脚
            logger.info("Processing footers")
            for section in doc.sections:
                footer = section.footer
                if footer:
                    for paragraph in footer.paragraphs:
                        if paragraph.text.strip():
                            text_parts.append(f"[Footer] {paragraph.text.strip()}")
            
            # 合并所有文本
            final_text = '\n\n'.join(text_parts)
            
            if not final_text.strip():
                logger.warning("No text content found in the document")
                return b"[Empty document or no text content found]"
            
            logger.info("Successfully extracted text from document")
            return final_text.encode('utf-8')
            
        except Exception as e:
            logger.error(f"Error converting Word document to text: {str(e)}")
            raise ValueError(f"Failed to convert Word document to text: {str(e)}")

class ExcelHandler(FormatHandler):
    def convert_to_pdf(self, input_data: bytes, dpi: int = 300) -> bytes:
        with NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_in, \
             NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_out:
            try:
                tmp_in.write(input_data)
                tmp_in.flush()
                
                # Use libreoffice for conversion
                cmd = f'soffice --headless --convert-to pdf "{tmp_in.name}" --outdir "{os.path.dirname(tmp_out.name)}"'
                logger.info(f"Executing command: {cmd}")
                result = os.system(cmd)
                
                if result != 0:
                    raise ValueError(f"LibreOffice conversion failed with exit code {result}")
                
                with open(tmp_out.name, 'rb') as f:
                    pdf_data = f.read()
                    
                return pdf_data
                
            finally:
                try:
                    os.unlink(tmp_in.name)
                    os.unlink(tmp_out.name)
                except Exception as e:
                    logger.error(f"Error cleaning up temporary files: {str(e)}")
            
    def convert_to_images(self, input_data: bytes, dpi: int = 300) -> List[Tuple[str, bytes]]:
        # First convert to PDF
        pdf_data = self.convert_to_pdf(input_data)
        
        # Then convert PDF pages to images
        pdf_handler = PDFHandler()
        return pdf_handler.convert_to_images(pdf_data, dpi)
        
    def convert_to_text(self, input_data: bytes) -> bytes:
        try:
            # Read Excel file
            df = pd.read_excel(io.BytesIO(input_data))
            
            # Convert to string representation
            text = df.to_string()
            if not text.strip():
                return b"[Empty spreadsheet]"
            return text.encode('utf-8')
        except Exception as e:
            logger.error(f"Error converting Excel to text: {str(e)}")
            raise ValueError(f"Failed to convert Excel to text: {str(e)}")

class PowerPointHandler(FormatHandler):
    def convert_to_pdf(self, input_data: bytes, dpi: int = 300) -> bytes:
        with NamedTemporaryFile(delete=False, suffix='.pptx') as tmp_in, \
             NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_out:
            try:
                tmp_in.write(input_data)
                tmp_in.flush()
                
                # Use libreoffice for conversion
                cmd = f'soffice --headless --convert-to pdf "{tmp_in.name}" --outdir "{os.path.dirname(tmp_out.name)}"'
                logger.info(f"Executing command: {cmd}")
                result = os.system(cmd)
                
                if result != 0:
                    raise ValueError(f"LibreOffice conversion failed with exit code {result}")
                
                with open(tmp_out.name, 'rb') as f:
                    pdf_data = f.read()
                    
                return pdf_data
                
            finally:
                try:
                    os.unlink(tmp_in.name)
                    os.unlink(tmp_out.name)
                except Exception as e:
                    logger.error(f"Error cleaning up temporary files: {str(e)}")
            
    def convert_to_images(self, input_data: bytes, dpi: int = 300) -> List[Tuple[str, bytes]]:
        pdf_data = self.convert_to_pdf(input_data)
        pdf_handler = PDFHandler()
        return pdf_handler.convert_to_images(pdf_data, dpi)
        
    def convert_to_text(self, input_data: bytes) -> bytes:
        try:
            prs = Presentation(io.BytesIO(input_data))
            text = []
            
            for i, slide in enumerate(prs.slides, 1):
                slide_text = []
                slide_text.append(f"\n[Slide {i}]")
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                
                if len(slide_text) > 1:  # 如果有实际内容（不只是标题）
                    text.extend(slide_text)
            
            final_text = '\n'.join(text)
            if not final_text.strip():
                return b"[Empty presentation or no text content found]"
                
            return final_text.encode('utf-8')
            
        except Exception as e:
            logger.error(f"Error converting PowerPoint to text: {str(e)}")
            raise ValueError(f"Failed to convert PowerPoint to text: {str(e)}")

def get_handler(format_str: str) -> FormatHandler:
    """Get the appropriate handler for a given format"""
    format_handlers = {
        # Word formats
        'doc': WordHandler(),
        'docx': WordHandler(),
        'wps': WordHandler(),
        # Excel formats
        'xls': ExcelHandler(),
        'xlsx': ExcelHandler(),
        'csv': ExcelHandler(),
        # PowerPoint formats
        'ppt': PowerPointHandler(),
        'pptx': PowerPointHandler(),
        # PDF format
        'pdf': PDFHandler(),
    }
    
    handler = format_handlers.get(format_str.lower())
    if not handler:
        raise ValueError(f"Unsupported format: {format_str}")
    
    return handler
