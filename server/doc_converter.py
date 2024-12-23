import os
import io
import logging
from enum import Enum
from typing import List, Optional, Tuple
import subprocess
import tempfile
from fastapi import HTTPException
from PIL import Image
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import pandas as pd
from b64encoder_decoder import custom_b64decode

# 设置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class SourceFormat(str, Enum):
    # Word formats
    DOC = "doc"
    DOCX = "docx"
    WPS = "wps"
    WPSS = "wpss"
    DOCM = "docm"
    DOTM = "dotm"
    DOT = "dot"
    DOTX = "dotx"
    HTML = "html"
    
    # PPT formats
    PPTX = "pptx"
    PPT = "ppt"
    POT = "pot"
    POTX = "potx"
    PPS = "pps"
    PPSX = "ppsx"
    DPS = "dps"
    DPT = "dpt"
    PPTM = "pptm"
    POTM = "potm"
    PPSM = "ppsm"
    DPSS = "dpss"
    
    # Excel formats
    XLS = "xls"
    XLT = "xlt"
    ET = "et"
    ETT = "ett"
    XLSX = "xlsx"
    XLTX = "xltx"
    CSV = "csv"
    XLSB = "xlsb"
    XLSM = "xlsm"
    XLTM = "xltm"
    ETS = "ets"
    
    # PDF format
    PDF = "pdf"

class TargetFormat(str, Enum):
    PDF = "pdf"
    PNG = "png"
    JPG = "jpg"
    TXT = "txt"

def parse_pages_param(pages_str: Optional[str]) -> List[int]:
    """
    Parse pages parameter from base64 encoded string
    Example: "1,2,4-10" -> [1, 2, 4, 5, 6, 7, 8, 9, 10]
    """
    if not pages_str:
        return []
        
    decoded = custom_b64decode(pages_str)
    pages = []
    
    for part in decoded.split(','):
        if '-' in part:
            start, end = map(int, part.split('-'))
            pages.extend(range(start, end + 1))
        else:
            pages.append(int(part))
            
    return sorted(list(set(pages)))

def convert_to_pdf(input_path: str, output_path: str, source_format: SourceFormat):
    """Convert document to PDF using LibreOffice"""
    try:
        logger.info(f"Converting {source_format} to PDF: {input_path} -> {output_path}")
        
        # 对于PDF文件，直接复制
        if source_format == SourceFormat.PDF:
            with open(input_path, 'rb') as src, open(output_path, 'wb') as dst:
                dst.write(src.read())
            return

        # 使用LibreOffice进行转换
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', os.path.dirname(output_path),
            input_path
        ]
        
        logger.info(f"Executing command: {' '.join(cmd)}")
        result = subprocess.run(cmd, capture_output=True, text=True)
        
        if result.returncode != 0:
            error_msg = f"LibreOffice conversion failed: {result.stderr}"
            logger.error(error_msg)
            raise HTTPException(status_code=500, detail=error_msg)
        
        # 重命名输出文件
        temp_pdf = os.path.join(
            os.path.dirname(output_path),
            os.path.splitext(os.path.basename(input_path))[0] + '.pdf'
        )
        os.rename(temp_pdf, output_path)
        logger.info("PDF conversion successful")
        
    except subprocess.CalledProcessError as e:
        error_msg = f"Failed to convert to PDF: {str(e)}"
        logger.error(error_msg)
        raise HTTPException(status_code=500, detail=error_msg)
    except Exception as e:
        error_msg = f"Unexpected error during PDF conversion: {str(e)}"
        logger.error(error_msg)
        raise HTTPException(status_code=500, detail=error_msg)

def convert_pdf_to_images(pdf_path: str, output_path: str, pages: List[int] = None, format: str = 'PNG', dpi: int = 300):
    """Convert PDF pages to images with specified DPI"""
    try:
        logger.info(f"Converting PDF to {format} images: {pdf_path} -> {output_path}")
        pdf_document = fitz.open(pdf_path)
        images = []
        
        # If no pages specified, convert all pages
        if not pages:
            pages = list(range(1, pdf_document.page_count + 1))
        
        logger.info(f"Processing pages: {pages}")
        for page_num in pages:
            if page_num > pdf_document.page_count:
                logger.warning(f"Skipping page {page_num} as it exceeds document length")
                continue
            
            page = pdf_document[page_num - 1]
            
            # 计算合适的缩放因子
            zoom = dpi/72
            mat = fitz.Matrix(zoom, zoom)
            
            try:
                # 获取页面图像
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                images.append(img)
                
                # 清理
                del pix
                
            except Exception as e:
                logger.error(f"Error processing page {page_num}: {str(e)}")
                raise ValueError(f"Failed to convert page {page_num} to image: {str(e)}")
        
        if not images:
            error_msg = "No valid pages to convert"
            logger.error(error_msg)
            raise HTTPException(status_code=400, detail=error_msg)
        
        # 保存图像
        if format.upper() == 'PDF':
            logger.info("Saving images as PDF")
            images[0].save(
                output_path,
                "PDF",
                save_all=True,
                append_images=images[1:] if len(images) > 1 else [],
                resolution=dpi
            )
        else:
            if len(images) > 1:
                logger.info(f"Saving multiple pages as {format}")
                images[0].save(
                    output_path,
                    format=format.upper(),
                    save_all=True,
                    append_images=images[1:],
                    optimize=True,
                    quality=95 if format.upper() in ['JPG', 'JPEG'] else None
                )
            else:
                logger.info(f"Saving single page as {format}")
                images[0].save(
                    output_path,
                    format=format.upper(),
                    optimize=True,
                    quality=95 if format.upper() in ['JPG', 'JPEG'] else None
                )
        
        logger.info("Image conversion successful")
        
    except Exception as e:
        error_msg = f"Failed to convert PDF to images: {str(e)}"
        logger.error(error_msg)
        raise HTTPException(status_code=500, detail=error_msg)
    finally:
        if 'pdf_document' in locals():
            pdf_document.close()

def extract_text_from_word(doc_path: str) -> str:
    """Extract text from Word document"""
    logger.info(f"Extracting text from Word document: {doc_path}")
    doc = Document(doc_path)
    text_parts = []
    
    # 处理文档主体
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())
    
    # 处理表格
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text_parts.append(' | '.join(row_text))
    
    return '\n\n'.join(text_parts)

def extract_text_from_ppt(ppt_path: str) -> str:
    """Extract text from PowerPoint document"""
    logger.info(f"Extracting text from PowerPoint document: {ppt_path}")
    prs = Presentation(ppt_path)
    text_parts = []
    
    for i, slide in enumerate(prs.slides, 1):
        slide_text = []
        slide_text.append(f"\n[Slide {i}]")
        
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        
        if len(slide_text) > 1:
            text_parts.extend(slide_text)
    
    return '\n'.join(text_parts)

def convert_pdf_to_text(pdf_path: str, output_path: str, pages: List[int] = None):
    """Extract text from PDF"""
    try:
        logger.info(f"Converting PDF to text: {pdf_path} -> {output_path}")
        pdf_document = fitz.open(pdf_path)
        text = []
        
        # If no pages specified, convert all pages
        if not pages:
            pages = list(range(1, pdf_document.page_count + 1))
        
        logger.info(f"Processing pages: {pages}")
        for page_num in pages:
            if page_num > pdf_document.page_count:
                logger.warning(f"Skipping page {page_num} as it exceeds document length")
                continue
            
            page = pdf_document[page_num - 1]
            page_text = page.get_text()
            if page_text.strip():  # 只添加非空页面
                text.append(f"[Page {page_num}]\n{page_text}")
        
        if text:
            logger.info("Writing extracted text to file")
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write('\n\n'.join(text))
        else:
            error_msg = "No valid pages to extract text from"
            logger.error(error_msg)
            raise HTTPException(status_code=400, detail=error_msg)
            
        logger.info("Text extraction successful")
        
    except Exception as e:
        error_msg = f"Failed to convert PDF to text: {str(e)}"
        logger.error(error_msg)
        raise HTTPException(status_code=500, detail=error_msg)
    finally:
        if 'pdf_document' in locals():
            pdf_document.close()

def convert_document(input_path: str, output_path: str, source_format: SourceFormat, 
                    target_format: TargetFormat, pages: List[int] = None, dpi: int = 300):
    """
    Convert document from source format to target format
    
    Args:
        input_path: Path to source document
        output_path: Path for converted document
        source_format: Source document format
        target_format: Target document format
        pages: Optional list of page numbers to convert
        dpi: Resolution for image conversion (default: 300)
        
    Supported conversions:
        - All formats -> PDF
        - All formats -> PNG/JPG (通过先转换为PDF再转换为图片)
        - Word/PPT/PDF -> TXT
    """
    logger.info(f"Converting document: {input_path} ({source_format}) -> {output_path} ({target_format})")
    
    # Create temporary directory for intermediate files
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            # 处理文本提取
            if target_format == TargetFormat.TXT:
                if source_format in [SourceFormat.DOC, SourceFormat.DOCX, SourceFormat.WPS]:
                    logger.info("直接从Word文档提取文本")
                    text = extract_text_from_word(input_path)
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(text)
                    return
                elif source_format in [SourceFormat.PPT, SourceFormat.PPTX, SourceFormat.DPS]:
                    logger.info("直接从PPT文档提取文本")
                    text = extract_text_from_ppt(input_path)
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(text)
                    return
            
            # 处理文档转换
            if target_format == TargetFormat.PDF:
                if source_format == SourceFormat.PDF and not pages:
                    logger.info("直接复制PDF文件")
                    with open(input_path, 'rb') as src, open(output_path, 'wb') as dst:
                        dst.write(src.read())
                else:
                    # 先转换为PDF（如果源格式不是PDF）
                    if source_format != SourceFormat.PDF:
                        logger.info(f"将{source_format}转换为PDF")
                        temp_pdf = os.path.join(temp_dir, 'temp.pdf')
                        convert_to_pdf(input_path, temp_pdf, source_format)
                        pdf_path = temp_pdf
                    else:
                        pdf_path = input_path

                    if pages:
                        logger.info("转换指定PDF页面")
                        convert_pdf_to_images(pdf_path, output_path, pages, 'PDF', dpi)
                    else:
                        logger.info("复制完整PDF文件")
                        with open(pdf_path, 'rb') as src, open(output_path, 'wb') as dst:
                            dst.write(src.read())
            
            elif target_format in [TargetFormat.PNG, TargetFormat.JPG]:
                # 先转换为PDF
                if source_format != SourceFormat.PDF:
                    logger.info(f"将{source_format}转换为PDF")
                    temp_pdf = os.path.join(temp_dir, 'temp.pdf')
                    convert_to_pdf(input_path, temp_pdf, source_format)
                    pdf_path = temp_pdf
                else:
                    pdf_path = input_path
                
                logger.info(f"将PDF转换为{target_format}图片")
                convert_pdf_to_images(pdf_path, output_path, pages, target_format, dpi)
            
            elif target_format == TargetFormat.TXT:
                if source_format in [SourceFormat.DOC, SourceFormat.DOCX, SourceFormat.WPS]:
                    logger.info("直接从Word文档提取文本")
                    text = extract_text_from_word(input_path)
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(text)
                elif source_format in [SourceFormat.PPT, SourceFormat.PPTX, SourceFormat.DPS]:
                    logger.info("直接从PPT文档提取文本")
                    text = extract_text_from_ppt(input_path)
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(text)
                else:
                    # 先转换为PDF再提取文本
                    if source_format != SourceFormat.PDF:
                        logger.info(f"将{source_format}转换为PDF")
                        temp_pdf = os.path.join(temp_dir, 'temp.pdf')
                        convert_to_pdf(input_path, temp_pdf, source_format)
                        pdf_path = temp_pdf
                    else:
                        pdf_path = input_path
                    
                    logger.info("从PDF提取文本")
                    convert_pdf_to_text(pdf_path, output_path, pages)
            else:
                error_msg = f"不支持的目标格式: {target_format}"
                logger.error(error_msg)
                raise HTTPException(status_code=400, detail=error_msg)
                
            logger.info("文档转换成功完成")
            
        except Exception as e:
            error_msg = f"文档转换失败: {str(e)}"
            logger.error(error_msg)
            raise HTTPException(status_code=500, detail=error_msg)
